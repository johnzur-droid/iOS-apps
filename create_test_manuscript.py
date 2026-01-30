#!/usr/bin/env python3
"""
Create a test manuscript document to verify the repair tool.

This creates a sample document with:
- Item headers (001 |, 002 |, ### 003 |, etc.)
- Related Topics with poisoned Google hyperlinks (#001, #002, etc.)
- Technical Sources with Google Search links (Search: AAP Policy)
- 🔝 emojis for navigation
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_hyperlink(paragraph, text, url):
    """Add an external hyperlink to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    # Blue color
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0000FF')
    rPr.append(color)

    # Underline
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    rPr.append(underline)

    new_run.append(rPr)

    text_elem = OxmlElement('w:t')
    text_elem.text = text
    new_run.append(text_elem)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

    return hyperlink


def create_test_document():
    """Create a test manuscript document."""
    doc = Document()

    # Title / TOC area
    title = doc.add_heading('The Zur Protocol Manuscript', 0)
    doc.add_paragraph('Table of Contents - Click 🔝 to return here')
    doc.add_paragraph()

    # Create several items
    items = [
        {
            'id': '001',
            'title': 'Introduction to the Protocol',
            'content': 'This section introduces the core concepts of the Zur Protocol.',
            'related': ['#002', '#003', '#019'],
            'sources': ['Search: AAP Policy', 'Search: Medical Guidelines 2024']
        },
        {
            'id': '002',
            'title': 'Safety Considerations',
            'content': 'Critical safety information for protocol implementation.',
            'related': ['#001', '#004', '#150'],
            'sources': ['Search: FDA Safety Standards']
        },
        {
            'id': '003',
            'title': 'Implementation Steps',
            'content': 'Step-by-step guide for implementing the protocol.',
            'related': ['#001', '#002'],
            'sources': ['Search: Implementation Best Practices']
        },
        {
            'id': '004',
            'title': 'Quality Assurance',
            'content': 'Quality control measures and verification procedures.',
            'related': ['#002', '#019'],
            'sources': ['Search: QA Standards ISO 9001']
        },
        {
            'id': '019',
            'title': 'Special Cases',
            'content': 'Handling edge cases and exceptions in the protocol.',
            'related': ['#001', '#004'],
            'sources': ['Search: Exception Handling Protocols']
        },
        {
            'id': '150',
            'title': 'Appendix A - Reference Tables',
            'content': 'Comprehensive reference tables for quick lookup.',
            'related': ['#001', '#002', '#003'],
            'sources': ['Search: Reference Data Tables']
        }
    ]

    for item in items:
        # Item Header (Object C)
        header_text = f"{item['id']} | {item['title']}"
        doc.add_heading(header_text, level=1)

        # Content
        doc.add_paragraph(item['content'])

        # Related Topics section with poisoned hyperlinks (Object A)
        related_para = doc.add_paragraph('Related Topics: ')
        for i, rel in enumerate(item['related']):
            if i > 0:
                related_para.add_run(', ')
            # Add poisoned Google hyperlink
            add_hyperlink(related_para, rel, f'https://www.google.com/search?q={rel[1:]}')

        # Technical Sources with good hyperlinks (Object B)
        sources_para = doc.add_paragraph('Technical Sources: ')
        for i, src in enumerate(item['sources']):
            if i > 0:
                sources_para.add_run(' | ')
            query = src.replace('Search: ', '').replace(' ', '+')
            add_hyperlink(sources_para, src, f'https://www.google.com/search?q={query}')

        # Navigation emoji
        doc.add_paragraph('🔝 Back to Top')
        doc.add_paragraph()  # Spacer

    # Also test "### XXX |" format
    doc.add_heading('### 042 | Alternative Header Format', level=1)
    doc.add_paragraph('This tests the alternative header format with ### prefix.')

    related_para = doc.add_paragraph('Related Topics: ')
    add_hyperlink(related_para, '#001', 'https://www.google.com/search?q=001')
    related_para.add_run(', ')
    add_hyperlink(related_para, '#042', 'https://www.google.com/search?q=042')

    doc.add_paragraph('🔝 Back to Top')

    # Save
    output_path = 'test_manuscript.docx'
    doc.save(output_path)
    print(f'✓ Created test document: {output_path}')
    print(f'  - {len(items) + 1} items with headers')
    print(f'  - Poisoned hyperlinks on #XXX references')
    print(f'  - Preserved hyperlinks on Search: references')
    print(f'  - 🔝 emojis for navigation')
    print()
    print('Now run: python manuscript_hyperlink_repair.py test_manuscript.docx')

    return output_path


if __name__ == '__main__':
    create_test_document()
